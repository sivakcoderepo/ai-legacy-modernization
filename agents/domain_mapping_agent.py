import env
from langchain_openai import ChatOpenAI
from typing import List, Dict
import json
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows

def create_comparison_excel(mapping_data: dict, output_path: str) -> str:
    """
    Creates a detailed Excel comparison report with multiple worksheets.
    
    Args:
        mapping_data: Domain mapping data from AI
        output_path: Path to save Excel file
    
    Returns:
        Path to the created Excel file
    """
    wb = Workbook()
    
    # Remove default sheet
    wb.remove(wb.active)
    
    # Define styles
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=12)
    highlight_fill = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
    success_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    warning_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
    error_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # 1. Summary Sheet
    ws_summary = wb.create_sheet("Summary")
    summary = mapping_data.get("mappingSummary", {})
    
    ws_summary.append(["Domain Model Comparison Report"])
    ws_summary.merge_cells('A1:B1')
    ws_summary['A1'].font = Font(bold=True, size=16)
    ws_summary.append([])
    
    ws_summary.append(["Metric", "Value"])
    ws_summary['A3'].fill = header_fill
    ws_summary['A3'].font = header_font
    ws_summary['B3'].fill = header_fill
    ws_summary['B3'].font = header_font
    
    ws_summary.append(["Total Source Entities", summary.get("totalSourceEntities", 0)])
    ws_summary.append(["Total Target Entities", summary.get("totalTargetEntities", 0)])
    ws_summary.append(["Direct Mappings", summary.get("directMappings", 0)])
    ws_summary.append(["Transformations Required", summary.get("transformationRequired", 0)])
    ws_summary.append(["Unmapped Source Entities", summary.get("unmappedSource", 0)])
    ws_summary.append(["Unmapped Target Entities", summary.get("unmappedTarget", 0)])
    ws_summary.append(["Overall Compatibility", summary.get("overallCompatibility", "Unknown")])
    
    # Color code compatibility
    compatibility_cell = ws_summary['B10']
    if summary.get("overallCompatibility") == "High":
        compatibility_cell.fill = success_fill
    elif summary.get("overallCompatibility") == "Medium":
        compatibility_cell.fill = warning_fill
    else:
        compatibility_cell.fill = error_fill
    
    ws_summary.column_dimensions['A'].width = 30
    ws_summary.column_dimensions['B'].width = 20
    
    # 2. Entity Mappings Sheet
    ws_entities = wb.create_sheet("Entity Mappings")
    entity_headers = ["Source Entity", "Target Entity", "Mapping Type", "Confidence %", "Complexity", "Notes"]
    ws_entities.append(entity_headers)
    
    for col_num, header in enumerate(entity_headers, 1):
        cell = ws_entities.cell(1, col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border
    
    for entity_map in mapping_data.get("entityMappings", []):
        ws_entities.append([
            entity_map.get("sourceEntity", ""),
            entity_map.get("targetEntity", ""),
            entity_map.get("mappingType", ""),
            entity_map.get("confidence", 0),
            entity_map.get("migrationComplexity", ""),
            entity_map.get("notes", "")
        ])
    
    # Color code by mapping type
    for row in range(2, ws_entities.max_row + 1):
        mapping_type = ws_entities.cell(row, 3).value
        if mapping_type == "Direct":
            ws_entities.cell(row, 3).fill = success_fill
        elif mapping_type == "Transformation":
            ws_entities.cell(row, 3).fill = warning_fill
        else:
            ws_entities.cell(row, 3).fill = error_fill
    
    # Auto-size columns
    for col in ws_entities.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2)
        ws_entities.column_dimensions[column].width = min(adjusted_width, 50)
    
    # 3. Property Mappings Sheet
    ws_properties = wb.create_sheet("Property Mappings")
    prop_headers = ["Entity", "Source Property", "Target Property", "Type Match", "Transformation", "Confidence %"]
    ws_properties.append(prop_headers)
    
    for col_num, header in enumerate(prop_headers, 1):
        cell = ws_properties.cell(1, col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
    
    for entity_map in mapping_data.get("entityMappings", []):
        entity_name = entity_map.get("sourceEntity", "")
        for prop_map in entity_map.get("propertyMappings", []):
            ws_properties.append([
                entity_name,
                prop_map.get("sourceProperty", ""),
                prop_map.get("targetProperty", ""),
                "Yes" if prop_map.get("dataTypeCompatible", False) else "No",
                prop_map.get("transformation", "None"),
                prop_map.get("confidence", 0)
            ])
    
    # Color code type compatibility
    for row in range(2, ws_properties.max_row + 1):
        type_match = ws_properties.cell(row, 4).value
        if type_match == "Yes":
            ws_properties.cell(row, 4).fill = success_fill
        else:
            ws_properties.cell(row, 4).fill = error_fill
    
    for col in ws_properties.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2)
        ws_properties.column_dimensions[column].width = min(adjusted_width, 50)
    
    # 4. Unmapped Items Sheet
    ws_unmapped = wb.create_sheet("Unmapped Items")
    ws_unmapped.append(["Type", "Entity/Property", "Reason", "Recommendation"])
    
    for col_num in range(1, 5):
        cell = ws_unmapped.cell(1, col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
    
    for unmapped in mapping_data.get("unmappedSourceEntities", []):
        ws_unmapped.append([
            "Source Entity",
            unmapped.get("entity", ""),
            unmapped.get("reason", ""),
            unmapped.get("recommendation", "")
        ])
    
    for unmapped in mapping_data.get("unmappedTargetEntities", []):
        ws_unmapped.append([
            "Target Entity",
            unmapped.get("entity", ""),
            unmapped.get("reason", ""),
            unmapped.get("dataSource", "")
        ])
    
    # Highlight unmapped items
    for row in range(2, ws_unmapped.max_row + 1):
        ws_unmapped.cell(row, 1).fill = error_fill
    
    for col in ws_unmapped.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2)
        ws_unmapped.column_dimensions[column].width = min(adjusted_width, 50)
    
    # 5. Transformation Rules Sheet
    ws_transforms = wb.create_sheet("Transformation Rules")
    ws_transforms.append(["Rule", "Source Pattern", "Target Pattern", "Complexity", "Example"])
    
    for col_num in range(1, 6):
        cell = ws_transforms.cell(1, col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
    
    for rule in mapping_data.get("transformationRules", []):
        ws_transforms.append([
            rule.get("rule", ""),
            rule.get("sourcePattern", ""),
            rule.get("targetPattern", ""),
            rule.get("complexity", ""),
            rule.get("example", "")
        ])
    
    for col in ws_transforms.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2)
        ws_transforms.column_dimensions[column].width = min(adjusted_width, 50)
    
    # 6. Data Quality Issues Sheet
    ws_quality = wb.create_sheet("Data Quality Issues")
    ws_quality.append(["Issue", "Impact", "Resolution"])
    
    for col_num in range(1, 4):
        cell = ws_quality.cell(1, col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = border
    
    for issue in mapping_data.get("dataQualityIssues", []):
        ws_quality.append([
            issue.get("issue", ""),
            issue.get("impact", ""),
            issue.get("resolution", "")
        ])
    
    for col in ws_quality.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2)
        ws_quality.column_dimensions[column].width = min(adjusted_width, 50)
    
    # 7. Recommendations Sheet
    ws_recommendations = wb.create_sheet("Recommendations")
    ws_recommendations.append(["Priority", "Recommendation"])
    
    ws_recommendations['A1'].fill = header_fill
    ws_recommendations['A1'].font = header_font
    ws_recommendations['B1'].fill = header_fill
    ws_recommendations['B1'].font = header_font
    
    for idx, rec in enumerate(mapping_data.get("recommendations", []), 1):
        ws_recommendations.append([idx, rec])
    
    ws_recommendations.column_dimensions['A'].width = 10
    ws_recommendations.column_dimensions['B'].width = 80
    
    # Save workbook
    wb.save(output_path)
    return output_path


def domain_mapping_agent(
    source_domain: str, 
    target_domain: str, 
    history: List[Dict[str, str]], 
    stream: bool = False
):
    """
    Maps source domain model to target domain model and generates Excel comparison report.
    
    Args:
        source_domain: Domain model extracted from VB code (JSON string)
        target_domain: Target domain model (JSON string)
        history: Conversation history
        stream: Whether to stream the response
    
    Returns:
        Excel file path and mapping analysis
    """
    llm = ChatOpenAI(model="gpt-4o", temperature=0)
    
    history_text = "\n".join([f'{h["role"]}: {h["content"]}' for h in history])
    
    # Try to parse as JSON if it's a string
    try:
        if isinstance(source_domain, str):
            source_domain_obj = json.loads(source_domain)
        else:
            source_domain_obj = source_domain
    except json.JSONDecodeError:
        source_domain_obj = {"raw": source_domain}
    
    try:
        if isinstance(target_domain, str):
            target_domain_obj = json.loads(target_domain)
        else:
            target_domain_obj = target_domain
    except json.JSONDecodeError:
        target_domain_obj = {"raw": target_domain}
    
    prompt = f"""
You are a domain modeling expert. Compare these two domain models and create a comprehensive mapping for migration.

Conversation history:
{history_text}

SOURCE DOMAIN MODEL (from legacy VB application):
{json.dumps(source_domain_obj, indent=2)}

TARGET DOMAIN MODEL (enterprise target architecture):
{json.dumps(target_domain_obj, indent=2)}

Analyze and create a JSON structure with the following format:
{{
  "mappingSummary": {{
    "totalSourceEntities": 0,
    "totalTargetEntities": 0,
    "directMappings": 0,
    "transformationRequired": 0,
    "unmappedSource": 0,
    "unmappedTarget": 0,
    "overallCompatibility": "High | Medium | Low"
  }},
  "entityMappings": [
    {{
      "sourceEntity": "VB Entity Name",
      "targetEntity": "Modern Entity Name",
      "confidence": 95,
      "mappingType": "Direct | Transformation | Split | Merge",
      "propertyMappings": [
        {{
          "sourceProperty": "vb_field",
          "targetProperty": "modernField",
          "dataTypeCompatible": true,
          "transformation": null or "description of needed transformation",
          "confidence": 90
        }}
      ],
      "notes": "Any special considerations",
      "migrationComplexity": "Low | Medium | High"
    }}
  ],
  "unmappedSourceEntities": [
    {{
      "entity": "VB Entity",
      "reason": "Why it couldn't be mapped",
      "recommendation": "What to do with it"
    }}
  ],
  "unmappedTargetEntities": [
    {{
      "entity": "Modern Entity",
      "reason": "No source equivalent",
      "dataSource": "How to populate this entity"
    }}
  ],
  "transformationRules": [
    {{
      "rule": "Description of transformation",
      "sourcePattern": "What needs to change",
      "targetPattern": "What it becomes",
      "complexity": "Low | Medium | High",
      "example": "Concrete example"
    }}
  ],
  "dataQualityIssues": [
    {{
      "issue": "Problem description",
      "impact": "How it affects migration",
      "resolution": "How to fix it"
    }}
  ],
  "recommendations": [
    "Prioritized recommendations for successful migration"
  ]
}}

Return ONLY the JSON structure, no markdown formatting or code blocks.
"""
    
    if stream:
        full_response = ""
        for chunk in llm.stream(prompt):
            chunk_text = chunk.content if hasattr(chunk, 'content') else str(chunk)
            full_response += chunk_text
            yield chunk
        
        # After streaming, generate Excel
        try:
            # Clean and parse response
            clean_response = full_response.replace("```json", "").replace("```", "").strip()
            mapping_data = json.loads(clean_response)
            
            # Generate Excel file
            excel_path = "/mnt/user-data/outputs/domain_mapping_comparison.xlsx"
            create_comparison_excel(mapping_data, excel_path)
            
            # Yield completion message with download link
            completion_msg = f"\n\n✅ Excel comparison report generated: {excel_path}\n"
            yield completion_msg
            
        except Exception as e:
            error_msg = f"\n\n⚠️ Could not generate Excel: {str(e)}\n"
            yield error_msg
    else:
        response = llm.invoke(prompt).content
        clean_response = response.replace("```json", "").replace("```", "").strip()
        
        try:
            mapping_data = json.loads(clean_response)
            excel_path = "/mnt/user-data/outputs/domain_mapping_comparison.xlsx"
            create_comparison_excel(mapping_data, excel_path)
            
            yield response + f"\n\n✅ Excel comparison report: {excel_path}"
        except Exception as e:
            yield response + f"\n\n⚠️ Could not generate Excel: {str(e)}"
